"""mixrace reporter 展示层测试|mixrace reporter presentation tests.

契约:TSV 机器可读输出逐字节不变;HTML 升级为完整文档+总览卡片+徽章+
排序/搜索交互+全部图中文图注;verdict 颜色永不单独出现(色+文字双编码)。
|Contract: TSV byte-identical; HTML upgraded (full doc, cards, badges,
sort/search, CN captions); verdict color never alone (color+label).
"""

import base64
import re
from pathlib import Path

import pytest

from biopytools.mixrace.reporter import (
    build_html_report,
    build_sample_report,
    build_summary_table,
    write_summary_excel,
)

# 任意字节即可:内嵌只读文件做 base64,不校验 PNG|any bytes: embed only b64s
_FAKE_PNG = b"\x89PNG\r\n\x1a\nfake-png-bytes"


def _row(**kw):
    """最小可判读行|minimal verdict row."""
    base = {"sample": "S1", "verdict": "pure", "subtag": "",
            "advice": "建议: 可保存|Advice: safe to keep",
            "rationale": "杂合率低于阈值|het below threshold",
            "het_rate": 0.0002, "robust_rate": 0.0001,
            "shared_only_rate": 0.0002, "top_partner": "S2",
            "mix_proportion": None, "het_rate_after_hotspot": 0.0003,
            "dp_ratio": 3.5, "host_rate": None, "pathogen_map_rate": 0.0,
            "contamination_rate": 0.0, "mean_depth": 12.3, "breadth_1x": 0.95}
    base.update(kw)
    return base


def _fig(tmp_path, name):
    """写一个假 PNG 并返回路径|write fake png, return path."""
    p = tmp_path / f"{name}.png"
    p.write_bytes(_FAKE_PNG)
    return str(p)


class TestSummaryTsvRegression:
    """TSV 回归守卫:重写展示层不得改变机器可读输出|TSV byte-identical guard."""

    def test_tsv_bytes_unchanged(self):
        tsv, _ = build_summary_table([_row()])
        expected = (
            "sample\tverdict\tadvice\thet_rate\trobust_rate\tshared_only_rate\t"
            "top_partner\tmix_proportion\thet_rate_after_hotspot\tdp_ratio\t"
            "host_rate\tpathogen_map_rate\tcontamination_rate\tmean_depth\t"
            "breadth_1x\n"
            "S1\tpure\t建议: 可保存|Advice: safe to keep\t0.0200%\t0.0100%\t"
            "0.0200%\tS2\t—\t0.0300%\t3.50\t—\t0.00%\t0.00%\t12.3x\t0.95%\n")
        assert tsv == expected

    def test_tsv_multi_row_keeps_input_order(self):
        rows = [_row(sample="S2", het_rate=0.5918, verdict="contaminated",
                     mix_proportion=0.88, dp_ratio=1.31),
                _row(sample="S1")]
        tsv, _ = build_summary_table(rows)
        lines = tsv.splitlines()
        assert len(lines) == 3
        assert lines[1].startswith("S2\tcontaminated")
        assert "\t59.18%\t" in lines[1] and "\t88%\t" in lines[1]


class TestStandaloneSummaryDocument:
    """verdict_summary.html 必须是完整文档(修 charset 乱码风险)|full doc fix."""

    def test_full_document_skeleton(self):
        _, html = build_summary_table([_row()])
        assert html.lstrip().startswith("<!DOCTYPE html>")
        assert "<meta charset=\"utf-8\">" in html
        assert "<title>" in html and "</html>" in html

    def test_contains_table_with_cn_headers(self):
        _, html = build_summary_table([_row()])
        assert "纯菌" in html
        assert "总杂合率" in html


class TestHtmlReportDocument:
    """主报告文档骨架|main report document skeleton."""

    def test_charset_and_title(self):
        html = build_html_report("标题T", [_row()], {})
        assert html.lstrip().startswith("<!DOCTYPE html>")
        assert "<meta charset=\"utf-8\">" in html
        assert "标题T" in html

    def test_escaping_sample_name(self):
        html = build_html_report("t", [_row(sample="<x>&s")], {})
        assert "<x>&s" not in html
        assert "&lt;x&gt;&amp;s" in html

    def test_verdict_note_included(self):
        html = build_html_report("t", [_row()], {}, verdict_note="判读口径: x<1%")
        assert "判读口径: x&lt;1%" in html


class TestOverviewCards:
    """总览卡片+构成条:verdict 计数一目了然|overview cards + composition strip."""

    def test_cards_count_verdicts(self):
        rows = [_row(sample=f"S{i}", verdict=v) for i, v in enumerate(
            ["pure", "pure", "contaminated", "divergent"])]
        html = build_html_report("t", rows, {})
        assert "纯菌" in html and "混杂菌株" in html
        assert "data-count=\"2\"" in html     # pure
        assert "data-count=\"1\"" in html     # divergent / contaminated
        assert "data-count=\"0\"" in html     # uncertain=0 也显示

    def test_composition_strip_segments_nonzero_only(self):
        rows = [_row(sample="S1", verdict="pure"),
                _row(sample="S2", verdict="contaminated")]
        html = build_html_report("t", rows, {})
        assert html.count('<div class="compo-seg') == 2

    def test_color_never_alone_segments_labeled(self):
        """色+文字双编码:每个构成条段落带 aria 标签|every segment labeled."""
        rows = [_row(sample="S1", verdict="pure"),
                _row(sample="S2", verdict="contaminated")]
        html = build_html_report("t", rows, {})
        segs = re.findall(r"<div[^>]*compo-seg[^>]*>", html)
        assert len(segs) == 2
        assert all("aria-label" in s for s in segs)
        assert any("纯菌" in s for s in segs)
        assert any("混杂" in s for s in segs)


class TestSummaryTablePresentation:
    """汇总表:横滚容器+粘性表头+data-v 排序属性+徽章|table presentation."""

    def test_scroll_wrapper_and_sticky_css(self):
        html = build_html_report("t", [_row()], {})
        assert "tbl-wrap" in html
        assert "position:sticky" in html or "position: sticky" in html

    def test_numeric_cells_carry_raw_data_v(self):
        html = build_html_report("t", [_row()], {})
        assert "data-v=\"0.0002\"" in html    # het_rate 原始值
        assert "data-v=\"3.5\"" in html       # dp_ratio

    def test_sortable_headers(self):
        html = build_html_report("t", [_row()], {})
        assert "data-sortable" in html

    def test_verdict_badge_with_dot_and_text(self):
        """徽章=色点+文字,永不只靠颜色|badge = dot + text, never color alone."""
        html = build_html_report("t", [_row(verdict="contaminated")], {})
        assert "badge" in html
        assert "混杂菌株" in html
        assert "dot" in html

    def test_subtag_shown_in_badge(self):
        html = build_html_report(
            "t", [_row(verdict="divergent", subtag="mild")], {})
        assert "轻度" in html


class TestFigureGallery:
    """图版:全部图+中文标题+图注|figure gallery: all figures + CN captions."""

    def test_figures_beyond_key_five_embedded(self, tmp_path):
        """dist_heatmap/altfrac_dist 不在旧 _KEY_FIGURES 也须入报告|beyond key5."""
        figs = {"pca_3view": _fig(tmp_path, "pca"),
                "dist_heatmap": _fig(tmp_path, "dist"),
                "altfrac_dist": _fig(tmp_path, "altfrac")}
        html = build_html_report("t", [_row()], figs)
        assert html.count("data:image/png;base64,") == 3

    def test_cn_title_and_caption_for_known_stem(self, tmp_path):
        html = build_html_report("t", [_row()],
                                 {"pca_3view": _fig(tmp_path, "pca")})
        assert "PCA" in html and "三视图" in html        # 中文标题
        assert "绿" in html and "红" in html             # 图注含判读色说明

    def test_unknown_stem_still_embedded_with_fallback_title(self, tmp_path):
        html = build_html_report("t", [_row()],
                                 {"custom_fig": _fig(tmp_path, "cust")})
        assert html.count("data:image/png;base64,") == 1
        assert "custom_fig" in html                      # 回退显示 stem

    def test_figure_order_follows_reading_logic(self, tmp_path):
        """按阅读逻辑排序:热图在前,三面板殿后|fixed reading order."""
        figs = {"eval_3panel": _fig(tmp_path, "e3"),
                "het_heatmap_100kb": _fig(tmp_path, "hh")}
        html = build_html_report("t", [_row()], figs)
        assert html.index("100kb") < html.index("三面板")

    def test_missing_figure_file_does_not_crash(self):
        html = build_html_report("t", [_row()],
                                 {"pca_3view": "/nonexistent/x.png"})
        assert "data:image/png;base64," not in html


class TestEvidenceSection:
    """逐样品证据链:中文化指标+徽章+搜索过滤|evidence: CN metrics + search."""

    def test_evidence_uses_cn_metric_names(self):
        html = build_html_report("t", [_row()], {})
        assert "逐样品证据链" in html
        assert "总杂合率" in html and "混合伴侣" in html

    def test_search_box_present(self):
        html = build_html_report("t", [_row()], {})
        assert "sample-search" in html

    def test_expand_all_buttons(self):
        html = build_html_report("t", [_row()], {})
        assert "全部展开" in html and "全部收起" in html


class TestInteractivityAndPrint:
    """原生 JS + 打印样式|vanilla JS + print CSS."""

    def test_sort_and_search_js_present(self):
        html = build_html_report("t", [_row()], {})
        assert "<script>" in html
        assert "sortTable" in html

    def test_sort_actually_wired_to_headers(self):
        """排序必须接线到表头(点击+键盘),不能只有函数定义|sort wired, not just defined."""
        html = build_html_report("t", [_row()], {})
        assert "document.querySelectorAll('th.sortable')" in html
        assert html.count("sortTable(th)") >= 3   # 定义+click+keydown|def+click+keydown

    def test_print_css(self):
        html = build_html_report("t", [_row()], {})
        assert "@media print" in html

    def test_no_emoji(self):
        """输出层不得引入 emoji(§5.2)|no emoji emitted by our layer."""
        html = build_html_report("t", [_row()], {})
        emoji_ranges = [
            (0x1F000, 0x1FAFF), (0x2600, 0x27BF), (0xFE00, 0xFE0F),
            (0x2B00, 0x2BFF), (0x1F1E6, 0x1F1FF)]
        for ch in html:
            if any(lo <= ord(ch) <= hi for lo, hi in emoji_ranges):
                pytest.fail(f"emoji found: {ch!r}")

    def test_no_hardcoded_abs_path(self):
        html = build_html_report("t", [_row()], {})
        assert "/share/org/" not in html


class TestMdAndExcelUntouched:
    """md 单样本报告与 Excel 不在本次范围,签名不变|md/excel signatures stable."""

    def test_md_report_still_works(self):
        md = build_sample_report("S1", _row(), {})
        assert "S1" in md and "混杂评估报告" in md

    def test_excel_still_writes(self, tmp_path):
        import openpyxl
        p = tmp_path / "s.xlsx"
        write_summary_excel([_row()], str(p))
        wb = openpyxl.load_workbook(str(p))
        assert "判读汇总" in wb.sheetnames and "summary" in wb.sheetnames
