"""genome2sv flank xlsx 写出测试|genome2sv flank xlsx writer tests"""

from openpyxl import load_workbook

from biopytools.genome2sv.utils import XLSX_MAX_CHARS, write_flank_xlsx


def make_row(sv_id="pan_sv.INS.00001", seq="ACGTACGTAC"):
    """单条记录构造器|single record builder"""
    return (sv_id, "c1", 200, 200, "INS", "10", 150, 250,
            [("1/1", "10", "60"), ("0/0", ".", ".")], seq)


def read_rows(path):
    """读回 xlsx 数据行|read back xlsx rows as lists"""
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = [tuple(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


class TestWriteFlankXlsx:
    def test_header_and_values(self, tmp_path):
        """表头与数据值正确|header and cell values"""
        out = tmp_path / "flank.xlsx"
        write_flank_xlsx([make_row()], ["S1", "S2"], str(out), "sv_flank300bp.fa")
        rows = read_rows(out)
        assert rows[0] == ("sv_id", "chrom", "pos", "end", "svtype", "svlen",
                           "flank_lo", "flank_hi",
                           "S1_GT", "S1_LN", "S1_QV", "S2_GT", "S2_LN", "S2_QV",
                           "sequence")
        assert rows[1] == ("pan_sv.INS.00001", "c1", 200, 200, "INS", "10",
                           150, 250, "1/1", "10", "60", "0/0", ".", ".",
                           "ACGTACGTAC")

    def test_oversize_sequence_placeholder(self, tmp_path):
        """超 Excel 单元格上限只写占位标记,全长见 fasta|oversize seq -> placeholder"""
        out = tmp_path / "flank.xlsx"
        big = "A" * (XLSX_MAX_CHARS + 1)
        write_flank_xlsx([make_row(sv_id="pan_sv.DEL.00002", seq=big)],
                         ["S1"], str(out), "sv_flank300bp.fa")
        cell = read_rows(out)[1][-1]
        assert len(str(cell)) < 200
        assert "pan_sv.DEL.00002" in str(cell)
        assert "sv_flank300bp.fa" in str(cell)

    def test_boundary_sequence_written(self, tmp_path):
        """恰好等于上限的序列正常写入|sequence exactly at limit is written"""
        out = tmp_path / "flank.xlsx"
        write_flank_xlsx([make_row(seq="C" * XLSX_MAX_CHARS)],
                         ["S1"], str(out), "f.fa")
        assert read_rows(out)[1][-1] == "C" * XLSX_MAX_CHARS

    def test_single_sample_columns(self, tmp_path):
        """样本数决定基因型列数|genotype columns follow sample count"""
        out = tmp_path / "flank.xlsx"
        write_flank_xlsx([make_row()], ["Only"], str(out), "f.fa")
        rows = read_rows(out)
        assert rows[0] == ("sv_id", "chrom", "pos", "end", "svtype", "svlen",
                           "flank_lo", "flank_hi",
                           "Only_GT", "Only_LN", "Only_QV", "sequence")
